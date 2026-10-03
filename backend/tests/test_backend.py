"""Run from the backend folder: python -m pytest"""

import io
import os
import tempfile
import zipfile

os.environ["WORTLAUF_DATA"] = tempfile.mkdtemp(prefix="wortlauf-test-")

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app import auth  # noqa: E402
from app.extract import ExtractError, extract  # noqa: E402
from app.main import app  # noqa: E402


def make_epub() -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as z:
        z.writestr("mimetype", "application/epub+zip")
        z.writestr(
            "META-INF/container.xml",
            '<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles>'
            '<rootfile full-path="OEBPS/content.opf"/></rootfiles></container>',
        )
        z.writestr(
            "OEBPS/content.opf",
            '<package xmlns="http://www.idpf.org/2007/opf" xmlns:dc="http://purl.org/dc/elements/1.1/">'
            "<metadata><dc:title>Der Testband</dc:title><dc:creator>A. Autorin</dc:creator></metadata>"
            '<manifest><item id="a" href="text/a.xhtml" media-type="application/xhtml+xml"/>'
            '<item id="b" href="text/b.xhtml" media-type="application/xhtml+xml"/>'
            '<item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/></manifest>'
            '<spine><itemref idref="a"/><itemref idref="b"/></spine></package>',
        )
        z.writestr(
            "OEBPS/toc.ncx",
            '<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/"><navMap>'
            '<navPoint><navLabel><text>Erstes Kapitel</text></navLabel><content src="text/a.xhtml"/></navPoint>'
            '<navPoint><navLabel><text>Zweites Kapitel</text></navLabel><content src="text/b.xhtml"/></navPoint>'
            '<navPoint><navLabel><text>Zweites, Teil B</text></navLabel><content src="text/b.xhtml#teil-b"/></navPoint>'
            "</navMap></ncx>",
        )
        z.writestr(
            "OEBPS/text/a.xhtml",
            "<html><head><title>x</title><style>p{}</style></head><body><h1>Erstes Kapitel</h1>"
            "<p>Ein Satz mit <em>Betonung</em> und&nbsp;Umbruch.</p><p>Zweiter Absatz.</p></body></html>",
        )
        z.writestr(
            "OEBPS/text/b.xhtml",
            '<html><body><h1>Zweites Kapitel</h1><p>Text.</p><h2 id="teil-b">Teil B</h2><p>Mehr Text.</p></body></html>',
        )
    return buffer.getvalue()


def make_pdf() -> bytes:
    from reportlab.lib.pagesizes import A5
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A5)
    for page in range(1, 5):
        pdf.drawString(40, 560, "Handbuch der Lagertechnik")
        y = 520
        lines = [
            f"Seite {page} beginnt mit einem langen Satz, der bis zum Rand der Zeile reicht und um-",
            "gebrochen wird, damit die Silbentrennung geprüft werden kann, und er endet",
            "hier kurz.",
            "Der zweite Absatz ist ebenfalls lang genug, um mehrere Zeilen zu füllen, und läuft",
            "über das Seitenende hinaus weiter, ohne Punkt am Ende der letzten Zeile",
        ]
        for line in lines:
            pdf.drawString(40, y, line)
            y -= 16
        pdf.drawString(200, 30, str(page))
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def test_epub_chapters_and_fragments():
    fmt, doc = extract("band.epub", make_epub())
    assert fmt == "epub"
    assert doc["title"] == "Der Testband" and doc["author"] == "A. Autorin"
    assert [c["title"] for c in doc["chapters"]] == ["Erstes Kapitel", "Zweites Kapitel", "Zweites, Teil B"]
    texts = [p["t"] for p in doc["paras"]]
    assert texts[1] == "Ein Satz mit Betonung und Umbruch."
    assert texts[doc["chapters"][2]["start"]] == "Teil B"
    assert all("ids" not in p for p in doc["paras"])


def test_pdf_paragraphs_headers_and_hyphens():
    fmt, doc = extract("handbuch.pdf", make_pdf())
    assert fmt == "pdf"
    text = " ".join(p["t"] for p in doc["paras"])
    assert "Handbuch der Lagertechnik" not in text  # running header removed
    assert "umgebrochen" in text  # hyphenated line break joined
    assert doc["paras"][0]["pg"] == 1
    assert doc["chapters"][0]["title"].startswith("Seite 1")
    # the paragraph left open at the page end must not swallow the next page's capitalised start
    assert any(p["t"].startswith("Seite 2 beginnt") for p in doc["paras"])


def test_markdown_and_text():
    fmt, doc = extract("notiz.md", "# Titel\n\nEin **fetter** [Link](http://x) Text.\n\n## Zwei\n\n- Punkt eins\n- Punkt zwei\n".encode())
    assert fmt == "md"
    assert [c["title"] for c in doc["chapters"]] == ["Titel", "Zwei"]
    assert doc["paras"][1]["t"] == "Ein fetter Link Text."
    assert [p["t"] for p in doc["paras"][-2:]] == ["Punkt eins", "Punkt zwei"]

    fmt, doc = extract("roman.txt", "Kapitel 1\n\nErste Zeile\nzweite Zeile.\n\nKapitel 2\n\nSchluss.".encode("cp1252"))
    assert fmt == "txt"
    assert [c["title"] for c in doc["chapters"]] == ["Kapitel 1", "Kapitel 2"]
    assert doc["paras"][1]["t"] == "Erste Zeile zweite Zeile."


def test_docx_and_html():
    import docx

    document = docx.Document()
    document.add_heading("Überschrift", level=1)
    document.add_paragraph("Absatz im Dokument.")
    buffer = io.BytesIO()
    document.save(buffer)
    fmt, doc = extract("brief.docx", buffer.getvalue())
    assert fmt == "docx" and doc["paras"][0].get("h") == 1 and doc["words"] == 4

    fmt, doc = extract("seite.html", b"<html><head><title>Seite</title><script>x()</script></head><body><h1>Kopf</h1><p>Inhalt &amp; mehr</p></body></html>")
    assert fmt == "html" and doc["title"] == "Seite"
    assert [p["t"] for p in doc["paras"]] == ["Kopf", "Inhalt & mehr"]


def test_rejects_unknown_and_empty():
    with pytest.raises(ExtractError):
        extract("bild.png", b"\x89PNG....")
    with pytest.raises(ExtractError):
        extract("leer.txt", b"   \n  ")


def test_api_flow():
    with TestClient(app) as client:
        assert client.get("/api/state").json() == {"setupNeeded": True, "authenticated": False, "user": None}
        assert client.get("/api/books").status_code == 401
        assert client.post("/api/setup", json={"name": "willi", "password": "kurz"}).status_code == 422
        assert client.post("/api/setup", json={"name": "willi", "password": "langes-passwort"}).status_code == 200
        assert client.post("/api/setup", json={"name": "x", "password": "langes-passwort"}).status_code == 409
        assert client.get("/api/state").json()["authenticated"] is True

        upload = client.post("/api/books", files={"file": ("band.epub", make_epub(), "application/epub+zip")})
        assert upload.status_code == 200, upload.text
        book = upload.json()["id"]
        assert client.post("/api/books", files={"file": ("x.png", b"\x89PNG", "image/png")}).status_code == 422
        assert len(client.get(f"/api/books/{book}/doc").json()["paras"]) == 7
        assert client.get(f"/api/books/{book}/file").content[:2] == b"PK"

        assert client.get(f"/api/books/{book}/progress").json() is None
        newer = {"para": 3, "offset": 4, "fraction": 0.5, "wpm": 320, "updated": 2000.0}
        older = {"para": 1, "offset": 0, "fraction": 0.1, "wpm": 300, "updated": 1000.0}
        assert client.put(f"/api/books/{book}/progress", json=newer).json()["para"] == 3
        assert client.put(f"/api/books/{book}/progress", json=older).json()["para"] == 3  # stale write ignored

        mark = client.post(f"/api/books/{book}/bookmarks", json={"para": 2, "offset": 0, "snippet": "Zweiter"}).json()
        assert len(client.get(f"/api/books/{book}/bookmarks").json()) == 1
        client.delete(f"/api/bookmarks/{mark['id']}")
        assert client.get(f"/api/books/{book}/bookmarks").json() == []

        client.post("/api/readings", json={"book_id": book, "started": __import__("time").time(), "active_ms": 60000, "words": 300})
        assert client.get("/api/stats").json()["total"]["words"] == 300

        client.put("/api/settings", json={"theme": "paper"})
        assert client.get("/api/settings").json() == {"theme": "paper"}

        plugins = client.get("/api/plugins").json()
        assert any(p["id"] == "builtin/layout-center" for p in plugins)
        assert client.get(plugins[0]["url"]).headers["content-type"].startswith("text/javascript")
        assert client.get("/api/plugins/local/..%2f..%2fbackend%2fapp%2fmain.py").status_code == 404

        assert client.post("/api/logout", headers={"origin": "https://evil.example"}).status_code == 403
        client.post("/api/logout")
        assert client.get("/api/books").status_code == 401
        for _ in range(auth.MAX_FAILS):
            assert client.post("/api/login", json={"name": "willi", "password": "falsch"}).status_code == 401
        assert client.post("/api/login", json={"name": "willi", "password": "langes-passwort"}).status_code == 429
        auth._fails["locked_until"] = 0
        assert client.post("/api/login", json={"name": "Willi", "password": "langes-passwort"}).status_code == 200
        assert client.delete(f"/api/books/{book}").status_code == 200
        assert client.get("/api/books").json() == []
